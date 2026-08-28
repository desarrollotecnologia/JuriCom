"""Verifica que la generación de minutas rellena los datos del contrato y
elimina los datos de los contratos de muestra usados como plantilla."""

import io
from types import SimpleNamespace

import docx
import pytest

from app.infrastructure.minutas.generator import PLANTILLAS, generar_minuta
from app.infrastructure.minutas.numeros import numero_a_letras


def _contrato():
    return SimpleNamespace(
        codigo="C-0042",
        proveedor_contratista="ACME INGENIERÍA S.A.S.",
        nit_proveedor="900.123.456-7",
        descripcion_servicio="mantenimiento correctivo de la caldera principal",
        valor=45_500_000,
        moneda="COP",
        plazo_cantidad=45,
        plazo_unidad="dias",
        forma_pago="Con anticipo del 30% y saldo contra entrega",
        tipo_precio="mas_iva",
    )


def _texto(data: bytes) -> str:
    doc = docx.Document(io.BytesIO(data))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                partes.append(cell.text)
    return "\n".join(partes)


MUESTRAS = {
    "obra": ["148.381.168", "FRANCISCO ANDRES RINCON SOLANO", "901.357.967", "2026-216"],
    "orden_trabajo": ["10.018.062", "VASQUEZ & RODRIGUEZ", "804.015.808-6", "2026-222"],
    "suministro": ["20.862.400", "INDUSTRIAS O&C", "901.358.312-9", "2026-213"],
    "cps": ["13.873.758", "METAL FULL", "METALFULL", "901.118.220", "2026-219"],
}


def test_numero_a_letras_muestras():
    assert numero_a_letras(45_500_000) == "CUARENTA Y CINCO MILLONES QUINIENTOS MIL"


@pytest.mark.parametrize("clave", list(PLANTILLAS))
def test_generar_minuta(clave):
    texto = _texto(generar_minuta(clave, _contrato()))
    assert "ACME INGENIERÍA S.A.S." in texto
    assert "900.123.456-7" in texto
    assert "C-0042" in texto
    assert "CUARENTA Y CINCO MILLONES QUINIENTOS MIL PESOS M/CTE" in texto
    assert "$45.500.000" in texto
    assert "CUARENTA Y CINCO (45) DÍAS" in texto
    assert "COMPLETAR" in texto
    for muestra in MUESTRAS[clave]:
        assert muestra not in texto, f"quedó dato de muestra: {muestra!r}"
