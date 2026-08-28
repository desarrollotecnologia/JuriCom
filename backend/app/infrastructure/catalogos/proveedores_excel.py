"""Lectura del catálogo de proveedores desde Excel.

Se usa tanto para el sembrado inicial (archivo empaquetado en `backend/data/`)
como para el botón "Importar Excel" del panel de Compras.

El Excel real trae 2 hojas (Servicios y Compras) con una fila de encabezados
alrededor de la fila 3 y datos debajo. Los encabezados se detectan por texto
(no por posición fija) para tolerar cambios menores de layout.
"""

from __future__ import annotations

import unicodedata
from io import BytesIO
from typing import BinaryIO, Union

import openpyxl


def normalizar_texto(valor: object) -> str:
    """minúsculas + sin acentos + espacios colapsados. Para comparar/buscar."""
    if valor is None:
        return ""
    s = str(valor)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _normalizar_para(valor: object) -> str:
    """El campo 'Para' viene desordenado (Servicios / Compras / Compras y
    Servicios / COMPRAS Y SERVICIO / ...). Lo reducimos a 3 categorías."""
    n = normalizar_texto(valor)
    tiene_serv = "servic" in n
    tiene_comp = "compra" in n
    if tiene_serv and tiene_comp:
        return "ambas"
    if tiene_serv:
        return "servicios"
    if tiene_comp:
        return "compras"
    return ""


# Mapa: campo destino -> lista de "pistas" que deben aparecer en el encabezado.
_CAMPOS = {
    "nombre": ["nombre del proveedor"],
    "sector": ["sector al que pertenec"],
    "para": ["para"],
    "ciudad": ["ciudad"],
    "telefono": ["telefono"],
    "correo": ["correo electr"],
    "contacto": ["contacto principal"],
    "condiciones_pago": ["condiciones de pago"],
    "calificacion": ["calificaci"],
}


def _detectar_encabezados(ws) -> tuple[int, dict[str, int]]:
    """Devuelve (fila_encabezado, {campo: columna}). Busca en las primeras filas
    la que contenga 'nombre del proveedor'."""
    max_scan = min(ws.max_row, 12)
    for r in range(1, max_scan + 1):
        textos = {
            c: normalizar_texto(ws.cell(r, c).value)
            for c in range(1, ws.max_column + 1)
        }
        if any("nombre del proveedor" in t for t in textos.values()):
            columnas: dict[str, int] = {}
            for campo, pistas in _CAMPOS.items():
                for c, t in textos.items():
                    if not t:
                        continue
                    if campo == "para":
                        # 'para' es una palabra corta; exigimos igualdad exacta
                        # para no confundir con "sector al que pertenece...".
                        if t == "para":
                            columnas[campo] = c
                            break
                    elif any(p in t for p in pistas):
                        columnas[campo] = c
                        break
            return r, columnas
    return 0, {}


def _valor(ws, fila: int, columnas: dict[str, int], campo: str) -> str:
    col = columnas.get(campo)
    if not col:
        return ""
    v = ws.cell(fila, col).value
    if v is None:
        return ""
    return " ".join(str(v).split())


def parse_catalogo_proveedores(
    source: Union[str, bytes, BinaryIO],
) -> list[dict]:
    """Lee todas las hojas y devuelve proveedores deduplicados por nombre.

    Si un proveedor aparece como Servicios en una hoja y Compras en otra, queda
    marcado como 'ambas'.
    """
    if isinstance(source, bytes):
        source = BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)

    por_nombre: dict[str, dict] = {}
    for ws in wb.worksheets:
        fila_enc, columnas = _detectar_encabezados(ws)
        if not fila_enc or "nombre" not in columnas:
            continue
        # Categoría por defecto según la hoja (SE=servicios, CO=compras).
        titulo = normalizar_texto(ws.title)
        default_para = "servicios" if titulo.endswith("se") else (
            "compras" if titulo.endswith("co") else ""
        )
        for r in range(fila_enc + 1, ws.max_row + 1):
            nombre = _valor(ws, r, columnas, "nombre")
            if not nombre:
                continue
            clave = normalizar_texto(nombre)
            if not clave:
                continue
            para = _normalizar_para(_valor(ws, r, columnas, "para")) or default_para
            telefono = _valor(ws, r, columnas, "telefono")
            if telefono.endswith(".0") and telefono[:-2].isdigit():
                telefono = telefono[:-2]  # openpyxl lee números como float
            datos = {
                "nombre": nombre,
                "sector": _valor(ws, r, columnas, "sector"),
                "para": para,
                "ciudad": _valor(ws, r, columnas, "ciudad"),
                "telefono": telefono,
                "correo": _valor(ws, r, columnas, "correo"),
                "contacto": _valor(ws, r, columnas, "contacto"),
                "condiciones_pago": _valor(ws, r, columnas, "condiciones_pago"),
                "calificacion": _valor(ws, r, columnas, "calificacion"),
            }
            existente = por_nombre.get(clave)
            if existente is None:
                por_nombre[clave] = datos
            else:
                # Fusiona: si difieren las categorías, es 'ambas'; rellena vacíos.
                if existente["para"] and para and existente["para"] != para:
                    existente["para"] = "ambas"
                elif not existente["para"]:
                    existente["para"] = para
                for k, v in datos.items():
                    if k != "para" and not existente.get(k) and v:
                        existente[k] = v

    return list(por_nombre.values())
