"""Genera un único Word (.docx) con todo el manual de usuario a partir de los .md.

Uso:  python build_docx.py

- Combina los capítulos en el orden definido en ORDEN.
- Convierte títulos, párrafos, listas, tablas, citas y código.
- Para cada imagen: si el archivo existe en img/ la incrusta; si no, deja un
  recuadro gris "[ Pega aquí la imagen: ... ]" para hacerlo en Word/Google Docs.

ponytail: mini-conversor Markdown->docx hecho a medida (sin pandoc, que no está
instalado). Techo: solo cubre la sintaxis usada en este manual (encabezados,
listas simples, tablas, citas, code fences, **negrita**/_cursiva_/`código`/[enlaces]).
Upgrade: instalar pandoc y usar `pandoc manual.md -o manual.docx`.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).parent
IMG = BASE / "img"

ORDEN = [
    "README.md",
    "trazabilidad-del-proceso.md",
    "rol-administrador.md",
    "rol-supervisor-solicitante.md",
    "rol-compras.md",
    "rol-proyectos.md",
    "rol-lider-aprobador.md",
    "rol-anticipos.md",
    "rol-contabilidad.md",
    "rol-tesoreria.md",
    "rol-juridica.md",
]

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|_[^_]+_)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMG_LINE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
SKIP_LINE = re.compile(r"^\[← Volver al índice\]")


def add_runs(paragraph, text):
    """Escribe `text` en el párrafo aplicando negrita/cursiva/código/enlaces."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(LINK.sub(r"\1", part[2:-2]))
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xB0, 0x2A, 0x37)
        elif part.startswith("_") and part.endswith("_"):
            run = paragraph.add_run(LINK.sub(r"\1", part[1:-1]))
            run.italic = True
        else:
            m = LINK.fullmatch(part)
            if m:
                paragraph.add_run(m.group(1))
            else:
                paragraph.add_run(part)


def shade(paragraph, color_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    pPr.append(shd)


def add_image(doc, path, alt, src):
    img_path = (path.parent / src).resolve()
    if img_path.exists():
        try:
            doc.add_picture(str(img_path), width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        except Exception:
            pass
    image_placeholder(doc, alt)


def image_placeholder(doc, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[  Pega aquí la imagen:  {alt}  ]")
    run.italic = True
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    shade(p, "EEF2F7")


def flush_table(doc, rows):
    # rows: lista de listas de celdas (texto). La 2ª fila es el separador ---.
    header = rows[0]
    data = [r for r in rows[2:]]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, cell in enumerate(header):
        c = table.rows[0].cells[i]
        c.paragraphs[0].text = ""
        add_runs(c.paragraphs[0], cell)
        for run in c.paragraphs[0].runs:
            run.bold = True
    for row in data:
        cells = table.add_row().cells
        for i in range(len(header)):
            cells[i].paragraphs[0].text = ""
            add_runs(cells[i].paragraphs[0], row[i] if i < len(row) else "")
    doc.add_paragraph()


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def convert_file(doc, path):
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_list_buffer = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if SKIP_LINE.match(stripped) or stripped == "---" or stripped == "":
            i += 1
            continue

        # Code fence
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # cierre
            if lang == "mermaid":
                p = doc.add_paragraph()
                r = p.add_run("» Diagrama de flujo (ver versión web del manual)")
                r.italic = True
                r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            else:
                for cl in code:
                    p = doc.add_paragraph()
                    r = p.add_run(cl)
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    shade(p, "F3F4F6")
            continue

        # Imagen (también si viene precedida de ">" dentro de una cita)
        mimg = IMG_LINE.match(stripped.lstrip(">").strip())
        if mimg:
            add_image(doc, path, mimg.group(1), mimg.group(2))
            i += 1
            continue

        # Tabla
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            if len(rows) >= 2:
                flush_table(doc, rows)
            continue

        # Encabezados
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
            i += 1
            continue

        # Cita / callout
        if stripped.startswith(">"):
            content = stripped.lstrip(">").strip()
            if content == "":
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            shade(p, "FFF7E6")
            add_runs(p, content)
            i += 1
            continue

        # Lista con viñeta
        m_ul = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        if m_ul:
            indent = len(m_ul.group(1))
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            p = doc.add_paragraph(style=style)
            add_runs(p, m_ul.group(2))
            i += 1
            continue

        # Lista numerada
        m_ol = re.match(r"^(\s*)\d+\.\s+(.*)$", raw)
        if m_ol:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m_ol.group(2))
            i += 1
            continue

        # Párrafo normal
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1


def main():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Portada
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Manual de Usuario")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(0x16, 0x39, 0x66)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run("JURICOM · Colbeef")
    rs.font.size = Pt(16)
    rs.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    doc.add_paragraph()

    for idx, name in enumerate(ORDEN):
        path = BASE / name
        if not path.exists():
            continue
        if idx > 0:
            doc.add_page_break()
        convert_file(doc, path)

    out = BASE / "Manual-de-Usuario-JURICOM.docx"
    doc.save(str(out))
    print(f"OK -> {out}")


if __name__ == "__main__":
    main()
